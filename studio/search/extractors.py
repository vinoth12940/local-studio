"""Search service: extractors. Moved verbatim from search.py (P6)."""
import subprocess
import tempfile
import threading
from pathlib import Path

from studio.search.chunking import MAX_CHUNKS, chunk_text

# ---------------------------------------------------------------- extractors
# extract() turns one file into items to embed: text chunks, images, or audio segments.

def _text_items(sections):
    """sections: [(locator, text)] -> text items, capped at MAX_CHUNKS for the whole file."""
    items = []
    for locator, text in sections:
        for chunk in chunk_text(text):
            items.append({"kind": "text", "text": chunk, "locator": locator})
            if len(items) == MAX_CHUNKS:
                return items
    return items


IMAGE_SIDE = 1024
MIN_IMAGE_SIDE = 64               # icons and spacers
SCAN_PAGES = 20
SAMPLE_RATE = 16000
SEGMENT_SECONDS = 30
MAX_SEGMENTS = 20                 # the first 10 minutes of a file
VIDEO_SECONDS = 600               # how much of a video is looked at
FRAME_SECONDS = 5                 # one frame this often: a hit points at the moment, and 10 minutes cost ~30 s to index


def _fit(im, side=IMAGE_SIDE):
    im = im.convert("RGB")
    im.thumbnail((side, side))
    return im


def _load_image(p, side=IMAGE_SIDE):
    from PIL import Image
    if p.suffix.lower() == ".heic":
        import pillow_heif
        pillow_heif.register_heif_opener()
    with Image.open(p) as im:
        if min(im.size) < MIN_IMAGE_SIDE:
            return None
        im.draft("RGB", (side, side))                     # lets the JPEG decoder downscale while decoding
        return _fit(im, side)


PDFIUM = threading.Lock()         # pdfium is not thread-safe: the indexer and the preview both render pages


def _pdf_items(p):
    from pypdf import PdfReader
    pages = [(pg.extract_text() or "").strip() for pg in PdfReader(str(p)).pages]
    chars = sum(len(t) for t in pages)
    if pages and chars < 80 * len(pages) and chars < 2000:            # scanned: same test as app.py's attachments
        import pypdfium2 as pdfium
        with PDFIUM:
            doc = pdfium.PdfDocument(str(p))
            return [{"kind": "image", "image": _fit(doc[i].render(scale=2).to_pil()), "locator": f"p{i + 1}"}
                    for i in range(min(SCAN_PAGES, len(doc)))]
    return _text_items([(f"p{i + 1}", t) for i, t in enumerate(pages) if t])


def _docx_text(p):
    import docx
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    for t in d.tables:
        parts.append("\n".join(" | ".join(c.text.strip() for c in r.cells) for r in t.rows))
    return "\n\n".join(parts)


def _xlsx_sections(p):
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    sections = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                rows.append(" | ".join("" if v is None else str(v) for v in row))
            if len(rows) >= 2000:
                break
        sections.append((ws.title, "\n".join(rows)))
    return sections


def _audio_segments(p):
    import soundfile as sf
    with tempfile.TemporaryDirectory() as d:
        wav = Path(d) / "a.wav"
        r = subprocess.run(["afconvert", "-f", "WAVE", "-d", f"LEI16@{SAMPLE_RATE}", "-c", "1", str(p), str(wav)],
                           capture_output=True, text=True, timeout=600)
        if r.returncode != 0:
            raise RuntimeError(f"afconvert failed: {r.stderr.strip()[:200]}")
        samples, _ = sf.read(wav, dtype="float32", frames=SAMPLE_RATE * SEGMENT_SECONDS * MAX_SEGMENTS)
    size = SAMPLE_RATE * SEGMENT_SECONDS
    segs = [samples[i:i + size] for i in range(0, len(samples), size)]
    return [s for i, s in enumerate(segs) if i == 0 or len(s) >= SAMPLE_RATE]     # drop a tail shorter than a second


def _frame_times(duration):
    end = min(duration, VIDEO_SECONDS)
    return [t for t in range(0, int(end) + 1, FRAME_SECONDS) if t < end] or [0]


def _video_frames(p, times=None, ahead=None):
    """[(seconds, image)] sampled from a video at `times` (default: every FRAME_SECONDS). `ahead`
    caps how many frames are decoded to reach one moment, for a thumbnail that must be quick."""
    import av
    out, shown = [], []
    with av.open(str(p)) as c:
        if not c.streams.video:
            return out
        st = c.streams.video[0]
        st.thread_type = "AUTO"
        tb = st.time_base
        duration = float(c.duration / av.time_base) if c.duration else float((st.duration or 0) * tb)
        want = list(_frame_times(duration) if times is None else times)

        def when(frame, t):
            return float(frame.pts * tb) if frame.pts is not None else float(t)

        def take(frame, t):
            at = when(frame, t)
            if not shown or abs(at - shown[-1]) >= 0.5:    # never the same frame twice
                shown.append(at)
                out.append((max(0, int(round(at))), _fit(frame.to_image())))

        try:
            # A seek lands on the keyframe before the time asked for. Many clips (screen recordings,
            # generated video) have one keyframe for the whole file, so check where a seek really lands:
            # if it is far from the target, read straight through once instead of seeking each time.
            dense = True
            if len(want) > 1:
                c.seek(int(want[1] / tb), stream=st)
                first = next(c.decode(st), None)
                dense = first is not None and when(first, want[1]) >= want[1] - 2.0
            if dense:
                for t in want:
                    c.seek(int(t / tb), stream=st)
                    for n, frame in enumerate(c.decode(st)):
                        if when(frame, t) >= t - 0.02 or (ahead and n >= ahead):
                            take(frame, t)
                            break
            else:
                c.seek(0, stream=st)
                i = 0
                for frame in c.decode(st):
                    if when(frame, want[i]) >= want[i] - 0.02:
                        take(frame, want[i])
                        i += 1
                        if i == len(want):
                            break
        except Exception:                                 # damaged part of the file: keep the frames we have
            pass
    return out


def extract(path, kind):
    p = Path(path)
    ext = p.suffix.lower()
    if kind == "video":
        items = [{"kind": "image", "image": im, "locator": f"{t}s"} for t, im in _video_frames(p)]
        try:
            sound = _audio_segments(p)
        except (RuntimeError, subprocess.TimeoutExpired):  # no soundtrack, or one afconvert can't read
            sound = []
        return items + [{"kind": "audio", "audio": seg, "locator": f"{i * SEGMENT_SECONDS}s"} for i, seg in enumerate(sound)]
    if kind == "image":
        im = _load_image(p)
        return [{"kind": "image", "image": im, "locator": ""}] if im is not None else []
    if kind == "audio":
        return [{"kind": "audio", "audio": seg, "locator": f"{i * SEGMENT_SECONDS}s"} for i, seg in enumerate(_audio_segments(p))]
    if ext == ".pdf":
        return _pdf_items(p)
    if ext == ".docx":
        return _text_items([("", _docx_text(p))])
    if ext in (".xlsx", ".xlsm"):
        return _text_items(_xlsx_sections(p))
    return _text_items([("", p.read_bytes().decode("utf-8", errors="replace"))])
